"""Read documents into pages of plain text (the "local" parser).

Supported: .pdf (page by page, tables kept as readable rows), .docx, .txt,
images (.png/.jpg), Excel (.xlsx/.csv) and PowerPoint (.pptx).
- Scanned PDF pages (a full-page image, or almost no text) are read with local
  OCR (rapidocr, see ocr.py), several pages in parallel. A scanned page's own
  text layer is ignored, because the OCR that made it is often poor ("I I I"
  instead of 111).
- Excel: one sheet = one page, each row as "Header: value; Header: value".
- PowerPoint: one slide = one page.
Word, text, CSV and image files have no real pages, so they are returned as
page 0 ("no page number") unless a .txt file contains form-feed page breaks.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from collections import Counter
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from .ocr import get_pool, ocr_image

log = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx", ".csv", ".pptx"} | IMAGE_EXTENSIONS
MIN_CHARS_PER_PAGE = 20  # fewer characters than this = "no text" (probably a scanned image)
SCANNED_IMAGE_COVERAGE = 0.8  # a page mostly covered by one image is treated as a scan
OCR_DPI = 200  # default page image resolution for OCR (setting OCR_DPI)
HEADER_SAMPLE_PAGES = 6  # while streaming, headers/footers are found from the first pages (one wave of OCR workers)

# progress(done, total, note): called after each page is read
Progress = Callable[[int, int, str], None]


@dataclass
class Page:
    number: int  # 1-based page number; 0 means "no page numbers" (docx/txt)
    text: str


@dataclass
class LoadedDocument:
    file_name: str
    pages: list[Page]
    page_count: int
    warnings: list[str] = field(default_factory=list)

    @property
    def has_text(self) -> bool:
        return any(p.text.strip() for p in self.pages)


class UnsupportedFileError(ValueError):
    pass


def load_document(path: Path, display_name: str | None = None, progress: Progress | None = None,
                  ocr_dpi: int = OCR_DPI, ocr_workers: int = 1) -> LoadedDocument:
    path = Path(path)
    name = display_name or path.name
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path, name, progress, ocr_dpi, ocr_workers)
    if ext == ".docx":
        return _load_docx(path, name)
    if ext == ".txt":
        return _load_txt(path, name)
    if ext in IMAGE_EXTENSIONS:
        return _load_image(path, name)
    if ext == ".xlsx":
        return _load_xlsx(path, name)
    if ext == ".csv":
        return _load_csv(path, name)
    if ext == ".pptx":
        return _load_pptx(path, name)
    raise UnsupportedFileError(
        f"{name}: file type '{ext}' is not supported (use PDF, Word, Excel, PowerPoint, text or image files).")


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _clean_cell(cell) -> str:
    return re.sub(r"\s+", " ", str(cell or "")).strip()


def table_to_text(rows: list[list]) -> str:
    """Turn a table into one line per row: 'Header: value; Header: value'."""
    rows = [[_clean_cell(c) for c in row] for row in rows if row and any(row)]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    if not body:
        return " | ".join(header)
    lines = []
    for row in body:
        parts = [f"{h}: {v}" if h else v for h, v in zip(header, row) if v]
        lines.append("; ".join(parts))
    return "\n".join(lines)


def _pdf_page_text(page) -> str:
    """Extract text in reading order, replacing each table with readable rows."""
    tables = sorted(page.find_tables(), key=lambda t: t.bbox[1])
    if not tables:
        return page.extract_text() or ""
    parts, top = [], 0.0
    for t in tables:
        x0, t_top, x1, t_bottom = t.bbox
        if t_top > top:
            above = page.within_bbox((0, top, page.width, t_top)).extract_text() or ""
            parts.append(above)
        parts.append(table_to_text(t.extract()))
        top = max(top, t_bottom)
    if top < page.height:
        parts.append(page.within_bbox((0, top, page.width, page.height)).extract_text() or "")
    return "\n".join(p for p in parts if p.strip())


def is_scanned_page(page) -> bool:
    """True if one image covers most of the page (a scan, maybe with a hidden OCR text layer)."""
    area = float(page.width * page.height) or 1.0
    for im in page.images:
        w, h = min(im["x1"], page.width) - max(im["x0"], 0), min(im["bottom"], page.height) - max(im["top"], 0)
        if w > 0 and h > 0 and w * h / area >= SCANNED_IMAGE_COVERAGE:
            return True
    return False


def _norm_line(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip().lower())


def repeated_lines(pages: list[Page]) -> set[str]:
    """Header/footer lines that repeat on most pages (e.g. 'Company | Policy', 'Page 3'), normalised."""
    if len(pages) < 3:
        return set()
    counts = Counter()
    for p in pages:
        counts.update({_norm_line(l) for l in p.text.splitlines() if l.strip()})
    return {line for line, n in counts.items() if n >= max(3, 0.6 * len(pages))}


def strip_lines(text: str, repeated: set[str]) -> str:
    return "\n".join(l for l in text.splitlines() if _norm_line(l) not in repeated) if repeated else text


def _strip_repeated_lines(pages: list[Page]) -> None:
    repeated = repeated_lines(pages)
    for p in pages:
        p.text = strip_lines(p.text, repeated)


@dataclass
class PdfPage:
    number: int
    total: int
    text: str
    ocr: bool  # True if the text was read with OCR


def iter_pdf_pages(path: Path, name: str, ocr_dpi: int = OCR_DPI, ocr_workers: int = 1,
                   known: dict[int, str] | None = None) -> Iterator[PdfPage]:
    """Yield the pages of a PDF in page order.

    Pages with a text layer are read directly. Scanned pages (or pages with almost no
    text) are rendered at ocr_dpi and OCR'd by a pool of ocr_workers processes, so
    several pages are read at once. `known` maps page numbers to text that was already
    read before (the parse cache); those pages are not read again.
    """
    import pdfplumber

    pool = get_pool(ocr_workers)
    known = known or {}
    window = 2 * pool.workers  # pages waiting for OCR at most (keeps memory bounded for big PDFs)
    try:
        pdf = pdfplumber.open(path)
    except Exception as exc:
        raise ValueError(f"{name}: could not be opened as a PDF (is it damaged or password-protected?)") from exc
    with pdf:
        total = len(pdf.pages)
        ready: dict[int, PdfPage] = {}
        pending: dict[int, tuple] = {}  # number -> (text layer, scanned?, image, future)

        def resolve(n: int) -> PdfPage:
            if n in ready:
                return ready.pop(n)
            text, scanned, image, future = pending.pop(n)
            try:
                ocr_text = future.result()
            except Exception:  # a worker process crashed: start a new pool, read this page here
                log.warning("OCR worker failed on %s page %d; reading it in the main process", name, n)
                pool.reset()
                ocr_text = ocr_image(image)
            if len(ocr_text.strip()) >= MIN_CHARS_PER_PAGE or scanned:
                return PdfPage(n, total, ocr_text, True)
            return PdfPage(n, total, text, False)

        next_page = 1
        try:
            for n in range(1, total + 1):
                if n in known:
                    ready[n] = PdfPage(n, total, known[n], False)
                else:
                    page = pdf.pages[n - 1]
                    try:
                        text = _pdf_page_text(page)
                    except Exception:  # fall back to plain text if table detection fails
                        log.warning("Table extraction failed on %s page %d; using plain text", name, n)
                        text = page.extract_text() or ""
                    # Glyphs without a text mapping (usually bullet symbols) come out as "(cid:N)".
                    text = re.sub(r"\(cid:\d+\)", "•", text)
                    scanned = is_scanned_page(page)
                    if scanned or len(text.strip()) < MIN_CHARS_PER_PAGE:
                        image = page.to_image(resolution=ocr_dpi).original
                        pending[n] = (text, scanned, image, pool.submit(image))
                    else:
                        ready[n] = PdfPage(n, total, text, False)
                    page.close()  # free pdfplumber's cached objects (big PDFs)
                # Hand out pages in order as soon as they are done; wait only if too many are queued.
                while next_page <= n and (next_page in ready or len(pending) >= window
                                          or pending[next_page][3].done()):
                    yield resolve(next_page)
                    next_page += 1
            while next_page <= total:
                yield resolve(next_page)
                next_page += 1
        finally:  # stopped early (e.g. the document was removed): drop pages still waiting for OCR
            for item in pending.values():
                item[3].cancel()


def _load_pdf(path: Path, name: str, progress: Progress | None = None, ocr_dpi: int = OCR_DPI,
              ocr_workers: int = 1) -> LoadedDocument:
    pages: list[Page] = []
    ocr_pages: list[int] = []
    try:
        for p in iter_pdf_pages(path, name, ocr_dpi, ocr_workers):
            pages.append(Page(p.number, p.text))
            if p.ocr:
                ocr_pages.append(p.number)
            if progress:
                progress(p.number, p.total, "")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"{name}: could not be opened as a PDF (is it damaged or password-protected?)") from exc

    _strip_repeated_lines(pages)
    doc = LoadedDocument(name, pages, page_count=len(pages))
    doc.warnings = pdf_warnings(name, pages, ocr_pages)
    doc.pages = [p for p in pages if len(p.text.strip()) >= MIN_CHARS_PER_PAGE]
    return doc


def pdf_warnings(name: str, pages: list[Page], ocr_pages: list[int]) -> list[str]:
    warnings = []
    empty = [p.number for p in pages if len(p.text.strip()) < MIN_CHARS_PER_PAGE]
    if pages and len(empty) == len(pages):
        warnings.append(
            f"{name} has no selectable text and OCR found none. It is probably a blank or "
            "unreadable scanned image, so it cannot be searched."
        )
    elif empty:
        warnings.append(
            f"{name}: page(s) {_ranges(empty)} have no selectable text "
            "(probably scanned images that OCR could not read) and were skipped."
        )
    read_by_ocr = [n for n in ocr_pages if n not in empty]
    if read_by_ocr:
        warnings.append(
            f"{name}: page(s) {_ranges(read_by_ocr)} are scanned and were read with OCR. "
            "Check important values against the original."
        )
    return warnings


def _ranges(numbers: list[int]) -> str:
    """[1, 2, 3, 7] -> '1-3, 7'"""
    out, start = [], None
    for i, n in enumerate(numbers):
        if start is None:
            start = n
        if i + 1 == len(numbers) or numbers[i + 1] != n + 1:
            out.append(str(start) if start == n else f"{start}-{n}")
            start = None
    return ", ".join(out)


# --------------------------------------------------------------------------- #
# Images (photos, screenshots)
# --------------------------------------------------------------------------- #
def _load_image(path: Path, name: str) -> LoadedDocument:
    from PIL import Image

    try:
        with Image.open(path) as im:
            im.load()
            text = ocr_image(im)
    except Exception as exc:
        raise ValueError(f"{name}: could not be opened as an image.") from exc
    doc = LoadedDocument(name, [Page(0, text)] if len(text.strip()) >= MIN_CHARS_PER_PAGE else [], page_count=1)
    if doc.has_text:
        doc.warnings.append(f"{name} was read with OCR. Check important values against the original.")
    else:
        doc.warnings.append(f"{name}: no readable text was found in this image.")
    return doc


# --------------------------------------------------------------------------- #
# DOCX / TXT
# --------------------------------------------------------------------------- #
def _load_docx(path: Path, name: str) -> LoadedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise ValueError(f"{name}: could not be opened as a Word document.") from exc

    blocks = []
    for child in d.element.body.iterchildren():  # paragraphs and tables in document order
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, d).text.strip()
            if text:
                blocks.append(text)
        elif tag == "tbl":
            rows = [[c.text for c in r.cells] for r in Table(child, d).rows]
            blocks.append(table_to_text(rows))
    text = "\n".join(b for b in blocks if b)
    doc = LoadedDocument(name, [Page(0, text)] if text.strip() else [], page_count=1)
    if not doc.has_text:
        doc.warnings.append(f"{name} contains no text.")
    return doc


def _decode(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def _load_txt(path: Path, name: str) -> LoadedDocument:
    text = _decode(path.read_bytes())
    if "\f" in text:  # form feeds mark page breaks
        parts = text.split("\f")
        pages = [Page(i, t) for i, t in enumerate(parts, start=1) if t.strip()]
        count = len(parts)
    else:
        pages = [Page(0, text)] if text.strip() else []
        count = 1
    doc = LoadedDocument(name, pages, page_count=count)
    if not doc.has_text:
        doc.warnings.append(f"{name} contains no text.")
    return doc


# --------------------------------------------------------------------------- #
# Excel / CSV
# --------------------------------------------------------------------------- #
def cell_text(value) -> str:
    """A cell value as text, without rounding (10.0 -> '10', 10.6 -> '10.6')."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _rows_text(rows: list[list]) -> str:
    rows = [[cell_text(v) for v in row] for row in rows]
    return table_to_text([r for r in rows if any(c.strip() for c in r)])


def _load_xlsx(path: Path, name: str) -> LoadedDocument:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"{name}: could not be opened as an Excel workbook.") from exc
    pages = []
    try:
        for i, ws in enumerate(wb.worksheets, start=1):
            body = _rows_text([list(r) for r in ws.iter_rows(values_only=True)])
            if body.strip():
                pages.append(Page(i, f"Sheet: {ws.title}\n{body}"))
        count = len(wb.worksheets)
    finally:
        wb.close()
    doc = LoadedDocument(name, pages, page_count=count)
    if not doc.has_text:
        doc.warnings.append(f"{name} contains no data.")
    return doc


def _load_csv(path: Path, name: str) -> LoadedDocument:
    text = _decode(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    body = _rows_text(list(csv.reader(io.StringIO(text), dialect)))
    doc = LoadedDocument(name, [Page(0, body)] if body.strip() else [], page_count=1)
    if not doc.has_text:
        doc.warnings.append(f"{name} contains no data.")
    return doc


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #
def _shape_texts(shapes, ocr_pictures: list) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out = []
    for shape in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            out += _shape_texts(shape.shapes, ocr_pictures)
        elif getattr(shape, "has_table", False) and shape.has_table:
            out.append(table_to_text([[c.text for c in row.cells] for row in shape.table.rows]))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
            out.append(shape.text_frame.text.strip())
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            ocr_pictures.append(shape)
    return out


def _load_pptx(path: Path, name: str) -> LoadedDocument:
    from PIL import Image
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ValueError(f"{name}: could not be opened as a PowerPoint file.") from exc
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        pictures: list = []
        parts = _shape_texts(slide.shapes, pictures)
        if len("".join(parts)) < MIN_CHARS_PER_PAGE:  # mostly pictures (e.g. a screenshot): OCR them
            for pic in pictures:
                try:
                    with Image.open(io.BytesIO(pic.image.blob)) as im:
                        parts.append(ocr_image(im))
                except Exception:
                    log.warning("Could not read a picture on %s slide %d", name, i)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        text = "\n".join(p for p in parts if p.strip())
        if text.strip():
            pages.append(Page(i, text))
    doc = LoadedDocument(name, pages, page_count=len(prs.slides))
    if not doc.has_text:
        doc.warnings.append(f"{name} contains no text.")
    return doc
