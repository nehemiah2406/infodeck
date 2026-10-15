import shutil

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from infodeck.loaders import load_document
from tests.conftest import SAMPLE_DIR

LEAVE = SAMPLE_DIR / "KRT_Leave_Policy.pdf"


def test_same_file_twice_is_not_duplicated(engine):
    first = engine.ingest_file(LEAVE)
    count = engine.store.count()
    second = engine.ingest_file(LEAVE)
    assert first.status == "added" and first.pages == 4 and first.chunks > 0
    assert second.status == "duplicate"
    assert engine.store.count() == count
    assert len(engine.store.list_documents()) == 1


def test_reupload_with_different_name_is_still_a_duplicate(engine):
    data = LEAVE.read_bytes()
    assert engine.save_and_ingest_upload("Leave Policy.pdf", data).status == "added"
    again = engine.save_and_ingest_upload("copy of leave.pdf", data)
    assert again.status == "duplicate"
    assert len(engine.store.list_documents()) == 1
    assert len(list(engine.settings.upload_dir.glob("*.pdf"))) == 1  # duplicate file not kept


def test_same_name_new_content_replaces_old_version(engine, tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("Version one says the office opens at 9am.")
    engine.ingest_file(f)
    f.write_text("Version two says the office opens at 8am.")
    result = engine.ingest_file(f)
    docs = engine.store.list_documents()
    assert result.status == "replaced"
    assert len(docs) == 1
    assert "8am" in engine.store.search("office opens", 1)[0].text


def test_remove_document_deletes_chunks_and_uploaded_file(engine):
    engine.save_and_ingest_upload("leave.pdf", LEAVE.read_bytes())
    doc = engine.store.list_documents()[0]
    engine.remove_document(doc.doc_id)
    assert engine.store.count() == 0
    assert not list(engine.settings.upload_dir.glob("*.pdf"))


def test_removing_a_sample_never_deletes_the_sample_file(engine):
    engine.ingest_file(LEAVE)
    engine.remove_document(engine.store.list_documents()[0].doc_id)
    assert LEAVE.exists()


def test_samples_load_only_on_first_start(engine):
    assert len(engine.load_samples_once()) == 3
    engine.remove_document(engine.store.list_documents()[0].doc_id)
    assert engine.load_samples_once() == []  # removed sample does not come back
    assert len(engine.store.list_documents()) == 2


def test_scanned_pdf_without_text_gives_clear_warning(engine, tmp_path):
    pdf = tmp_path / "scanned.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    for _ in range(2):  # pages with only a picture-like shape, no text
        c.rect(100, 100, 300, 500, fill=1)
        c.showPage()
    c.save()
    doc = load_document(pdf)
    assert not doc.has_text
    assert "no selectable text" in doc.warnings[0] and "scanned" in doc.warnings[0]
    result = engine.ingest_file(pdf)
    assert result.status == "no_text" and "scanned" in result.message
    assert engine.store.count() == 0


def test_txt_and_docx_are_supported(engine, tmp_path):
    import docx

    txt = tmp_path / "faq.txt"
    txt.write_text("The canteen is open from 12:00 to 15:00.", encoding="utf-8")
    d = docx.Document()
    d.add_paragraph("Parking passes are issued by the facilities team.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Floor", "Team"
    t.cell(1, 0).text, t.cell(1, 1).text = "3", "Finance"
    d.save(tmp_path / "office.docx")

    assert engine.ingest_file(txt).status == "added"
    assert engine.ingest_file(tmp_path / "office.docx").status == "added"
    assert "Floor: 3; Team: Finance" in load_document(tmp_path / "office.docx").pages[0].text
    hit = engine.store.search("canteen open", 1)[0]
    assert hit.file_name == "faq.txt" and hit.label == "faq.txt"  # no page number for .txt


def test_unsupported_and_damaged_files_return_friendly_errors(engine, tmp_path):
    bad = tmp_path / "image.png"
    bad.write_bytes(b"\x89PNG")
    assert engine.save_and_ingest_upload("image.png", b"\x89PNG").status == "error"
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"this is not a pdf")
    result = engine.ingest_file(broken)
    assert result.status == "error" and "could not be opened" in result.message


def test_upload_file_name_cannot_escape_upload_folder(engine):
    result = engine.save_and_ingest_upload(r"..\..\evil.pdf", LEAVE.read_bytes())
    assert result.status == "added"
    saved = list(engine.settings.upload_dir.glob("*.pdf"))
    assert len(saved) == 1 and saved[0].parent == engine.settings.upload_dir
