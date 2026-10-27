"""Complex documents: loaders, OCR, Gemini page parsing (fake client), cache, structure-aware
chunking, the exact-value index and the number check. No network; fictional data only."""

from __future__ import annotations

import io
import json
from dataclasses import replace
from types import SimpleNamespace

import pytest
from google.genai import errors

from infodeck.chunking import chunk_pages
from infodeck.engine import InfoDeckEngine, unverified_numbers
from infodeck.facts import extract_facts, question_covered
from infodeck.gemini_parser import GeminiPageReader, ParseError
from infodeck.loaders import Page, load_document
from infodeck.store import Passage, VectorStore
from scripts.make_eval_docs import make_faq_docx, make_scanned_pdf, make_torque_xlsx, render_scan
from tests.conftest import FakeEmbedder, FakeLLM, llm_status

CALC_LINES = [
    "9. Our proposal of push-up distance, (Pd), mm",
    "   (1) Push-up distance, 10.6 mm",
    "   (2) Grip stress corresponding to Pd, N/mm²   σra = Pd / Pr x σr = 111 N/mm²",
    "10. Variables",
    "   A : Interfacial surface area .................. A= 1,118,355",
]
CALC_TEXT = "\n".join(line.strip() for line in CALC_LINES)
GEMINI_MARKDOWN = CALC_TEXT  # what the fake Gemini "reads" from every page image


class FakeGeminiParser:
    """Stands in for google.genai.Client: records every request, replies with fixed Markdown."""

    def __init__(self, reply: str = GEMINI_MARKDOWN, fail: list[Exception] | None = None):
        self.models = self
        self.reply = reply
        self.fail = list(fail or [])
        self.calls: list[list] = []

    def generate_content(self, model, contents, config):
        self.calls.append(contents)
        if self.fail:
            raise self.fail.pop(0)
        return SimpleNamespace(text=self.reply)


def rate_limit_error(delay: str = "7s") -> errors.ClientError:
    return errors.ClientError(429, {"error": {"code": 429, "message": "Resource exhausted", "status":
                                              "RESOURCE_EXHAUSTED", "details": [{
                                                  "@type": "type.googleapis.com/google.rpc.RetryInfo",
                                                  "retryDelay": delay}]}})


@pytest.fixture
def scanned_pdf(tmp_path):
    """Two-page image-only 'scan' whose hidden text layer is bad OCR ('I I I', 'Ora', '(l)')."""
    path = tmp_path / "calc_scanned.pdf"
    make_scanned_pdf(path, [CALC_LINES, ["11. Hydraulic oil injection", "   (1) Required oil pressure, 142 N/mm²"]])
    return path


def gemini_engine(settings, client) -> InfoDeckEngine:
    s = replace(settings, parser="gemini")
    return InfoDeckEngine(VectorStore(s.chroma_dir, FakeEmbedder()), s, gemini_client=client)


# ---- loaders ---------------------------------------------------------------- #
def test_xlsx_rows_become_header_value_lines_per_sheet(tmp_path):
    make_torque_xlsx(tmp_path / "torque.xlsx")
    doc = load_document(tmp_path / "torque.xlsx")
    assert doc.page_count == 2 and [p.number for p in doc.pages] == [1, 2]
    assert doc.pages[0].text.startswith("Sheet: Torque")
    assert "Bolt size: M16; Property class: 8.8; Tightening torque (N·m): 195; Preload (kN): 70.6" in doc.pages[0].text
    assert "Check: Coupling alignment check; Interval (months): 6" in doc.pages[1].text


def test_csv_rows_become_header_value_lines(tmp_path):
    (tmp_path / "limits.csv").write_text("Part,Limit (mm),Note\nShaft,0.05,runout\nHub,0.10,flatness\n",
                                         encoding="utf-8")
    text = load_document(tmp_path / "limits.csv").pages[0].text
    assert "Part: Hub; Limit (mm): 0.10; Note: flatness" in text  # "0.10" kept exactly, not 0.1


def test_pptx_one_slide_is_one_page_with_tables_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Coupling overview"
    rows = s1.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for c, (h, v) in enumerate([("Item", "Push-up distance"), ("Value", "10.6 mm")]):
        rows.cell(0, c).text, rows.cell(1, c).text = h, v
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Next steps"
    s2.notes_slide.notes_text_frame.text = "Check the dial gauge before mounting."
    prs.save(tmp_path / "deck.pptx")

    doc = load_document(tmp_path / "deck.pptx")
    assert doc.page_count == 2 and [p.number for p in doc.pages] == [1, 2]
    assert "Item: Push-up distance; Value: 10.6 mm" in doc.pages[0].text
    assert "Notes: Check the dial gauge before mounting." in doc.pages[1].text


def test_docx_qa_list_is_read_in_order(tmp_path):
    make_faq_docx(tmp_path / "faq.docx", [("Who approves changes?", "The lead engineer approves changes.")])
    text = load_document(tmp_path / "faq.docx").pages[0].text
    assert "Q: Who approves changes?\nA: The lead engineer approves changes." in text


# ---- OCR (local, rapidocr) ------------------------------------------------- #
def test_image_file_is_read_with_ocr(tmp_path):
    render_scan(CALC_LINES[:2]).save(tmp_path / "photo.png")
    doc = load_document(tmp_path / "photo.png")
    assert doc.pages and doc.pages[0].number == 0
    assert "Push-up distance, 10.6 mm" in doc.pages[0].text
    assert "OCR" in doc.warnings[0]


def test_local_parser_ocrs_scans_instead_of_trusting_a_bad_text_layer(scanned_pdf):
    import pdfplumber

    with pdfplumber.open(scanned_pdf) as pdf:
        assert "I I I" in pdf.pages[0].extract_text()  # the bad layer really is there
    text = load_document(scanned_pdf).pages[0].text
    assert "I I I" not in text and "(l)" not in text
    assert "= 111 N/mm²" in text and "(1) Push-up distance, 10.6 mm" in text


# ---- Gemini page parser (fake client, no network) --------------------------- #
def test_gemini_mode_sends_page_images_never_the_text_layer(settings, scanned_pdf):
    from PIL import Image

    client = FakeGeminiParser()
    result = gemini_engine(settings, client).ingest_file(scanned_pdf)
    assert result.status == "added" and result.pages == 2
    assert len(client.calls) == 2  # one request per page
    for contents in client.calls:
        image, prompt = contents
        assert image.inline_data.mime_type == "image/png"
        with Image.open(io.BytesIO(image.inline_data.data)) as im:
            assert abs(im.width - 1654) <= 2  # A4 rendered at 200 dpi
        sent_text = prompt if isinstance(prompt, str) else ""
        assert "I I I" not in sent_text and "Ora" not in sent_text and "Push-up" not in sent_text


def test_gemini_rate_limit_429_waits_and_retries():
    waits, notes = [], []
    client = FakeGeminiParser(fail=[rate_limit_error("7s"), rate_limit_error("7s")])
    reader = GeminiPageReader("", "gemini-test", client=client, sleep=waits.append)
    text = reader.read(b"\x89PNG fake", on_wait=lambda delay, reason: notes.append(reason))
    assert text == GEMINI_MARKDOWN
    assert len(client.calls) == 3 and waits == [8.0, 8.0]  # Google's retryDelay + 1 s
    assert notes == ["Gemini rate limit reached"] * 2


def test_gemini_gives_up_after_max_attempts_with_a_friendly_message():
    client = FakeGeminiParser(fail=[rate_limit_error()] * 3)
    reader = GeminiPageReader("", "gemini-test", client=client, sleep=lambda s: None, max_attempts=3)
    with pytest.raises(ParseError, match="upload the file again later"):
        reader.read(b"png")


def test_used_up_daily_quota_stops_at_once_instead_of_retrying():
    daily = errors.ClientError(429, {"error": {"code": 429, "status": "RESOURCE_EXHAUSTED", "details": [
        {"violations": [{"quotaId": "GenerateRequestsPerDayPerProjectPerModel-FreeTier"}]},
        {"@type": "type.googleapis.com/google.rpc.RetryInfo", "retryDelay": "82863s"}]}})
    waits = []
    client = FakeGeminiParser(fail=[daily])
    with pytest.raises(ParseError, match="used up for today.*about 23 h"):
        GeminiPageReader("", "gemini-test", client=client, sleep=waits.append).read(b"png")
    assert len(client.calls) == 1 and waits == []


def test_gemini_progress_reports_page_x_of_n(settings, scanned_pdf):
    calls = []
    gemini_engine(settings, FakeGeminiParser()).ingest_file(scanned_pdf, progress=lambda d, t, n: calls.append((d, t)))
    assert calls[-1] == (2, 2) and all(t == 2 for _, t in calls)


def test_parsed_pages_are_saved_and_reused_on_reupload(settings, scanned_pdf):
    client = FakeGeminiParser()
    engine = gemini_engine(settings, client)
    data = scanned_pdf.read_bytes()
    first = engine.save_and_ingest_upload("calc.pdf", data)
    assert first.status == "added" and len(client.calls) == 2

    [doc] = engine.store.list_documents()
    folder = settings.parsed_dir / doc.doc_id
    assert (folder / "page_1.md").read_text(encoding="utf-8") == GEMINI_MARKDOWN
    meta = json.loads((folder / "meta.json").read_text(encoding="utf-8"))
    assert meta["file_name"] == "calc.pdf" and meta["parser"].startswith("gemini:") and meta["complete"]

    engine.remove_document(doc.doc_id)
    again = engine.save_and_ingest_upload("calc.pdf", data)
    assert again.status == "added" and len(client.calls) == 2  # cache hit: Gemini not called again


def test_interrupted_gemini_parse_continues_where_it_stopped(settings, scanned_pdf):
    broken = FakeGeminiParser(fail=[errors.ClientError(400, {"error": {"message": "bad page"}})] * 2)
    engine = gemini_engine(settings, broken)
    result = engine.ingest_file(scanned_pdf)
    assert result.status == "error" and "stopped at page" in result.message

    working = FakeGeminiParser()
    engine.gemini_client = working
    done_before = len(list(settings.parsed_dir.glob("*/page_*.md")))
    assert engine.ingest_file(scanned_pdf).status == "added"
    assert len(working.calls) == 2 - done_before  # only the missing pages are read


def test_local_parser_also_saves_page_files(settings, engine, tmp_path):
    (tmp_path / "notes.txt").write_text("Page one text here.\fPage two text here.", encoding="utf-8")
    engine.ingest_file(tmp_path / "notes.txt")
    [doc] = engine.store.list_documents()
    assert (settings.parsed_dir / doc.doc_id / "page_2.md").read_text(encoding="utf-8") == "Page two text here."


# ---- structure-aware chunking ---------------------------------------------- #
def test_numbered_item_stays_in_one_chunk_with_its_sub_items():
    chunks = chunk_pages([Page(3, "8. Fit data\n(1) Shaft diameter, 180 mm\n" + CALC_TEXT)], "calc.pdf", 300, 50)
    item9 = [c for c in chunks if "9. Our proposal" in c.text]
    assert len(item9) == 1 and "σra = Pd / Pr x σr = 111 N/mm²" in item9[0].text


def test_long_numbered_item_repeats_its_title_on_every_part():
    lines = [f"({i}) Check value number {i}, {i}.5 mm" for i in range(1, 60)]
    chunks = chunk_pages([Page(1, "9. Our proposal of push-up distance, (Pd), mm\n" + "\n".join(lines))],
                         "calc.pdf", 400, 80)
    assert len(chunks) > 3 and all(len(c.text) <= 400 for c in chunks)
    assert all(c.text.startswith("9. Our proposal of push-up distance") for c in chunks)
    assert all(f"({i}) Check value number {i}, {i}.5 mm" in "\n".join(c.text for c in chunks) for i in range(1, 60))


def test_item_continued_on_next_page_keeps_its_title():
    pages = [Page(1, "9. Our proposal of push-up distance, (Pd), mm\n(1) Push-up distance, 10.6 mm"),
             Page(2, "(2) Grip stress corresponding to Pd, N/mm²   σra = Pd / Pr x σr = 111 N/mm²")]
    page2 = [c for c in chunk_pages(pages, "calc.pdf") if c.page == 2]
    assert page2[0].text.startswith("9. Our proposal of push-up distance, (Pd), mm (continued)")


def test_split_markdown_table_repeats_its_header():
    rows = "\n".join(f"| M{i} | {i * 10} |" for i in range(10, 80))
    chunks = chunk_pages([Page(1, "1. Torque table\n| Bolt | Torque (N·m) |\n|---|---|\n" + rows)], "t.pdf", 300, 50)
    assert len(chunks) > 2
    assert all("| Bolt | Torque (N·m) |" in c.text for c in chunks)


def test_question_and_answer_stay_together():
    text = "Q: Who approves changes?\nA: The lead engineer.\nQ: How long are records kept?\nA: 10 years."
    chunks = chunk_pages([Page(0, text)], "faq.docx", 60, 10)
    assert [c.text for c in chunks] == ["Q: Who approves changes?\nA: The lead engineer.",
                                        "Q: How long are records kept?\nA: 10 years."]


# ---- exact-value / Q&A index ----------------------------------------------- #
def test_fact_extraction_keeps_values_exactly():
    facts = {f.name: f for f in extract_facts([Page(3, CALC_TEXT)], "calc.pdf")}
    pd = facts["Push-up distance"]
    assert (pd.value, pd.unit, pd.item, pd.page) == ("10.6", "mm", "9.(1)", 3)
    grip = facts["Grip stress corresponding to Pd"]
    assert (grip.symbol, grip.formula, grip.value, grip.unit) == ("σra", "σra = Pd / Pr x σr", "111", "N/mm²")
    area = facts["Interfacial surface area"]
    assert (area.symbol, area.value) == ("A", "1,118,355")  # thousands separators kept


def test_fact_extraction_tables_and_qa():
    text = ("Grade: S355; Yield strength (MPa): 355; Tensile strength (MPa): 470\n"
            "| Bolt size | Tightening torque (N·m) |\n|---|---|\n| M20 | 540 |\n\n"
            "Q: Who approves changes?\nA: The lead engineer approves changes.")
    facts = {f.name: f for f in extract_facts([Page(1, text)], "mixed.pdf")}
    assert (facts["Grade S355 Tensile strength"].value, facts["Grade S355 Tensile strength"].unit) == ("470", "MPa")
    assert facts["Bolt size M20 Tightening torque"].value == "540"
    assert facts["Who approves changes?"].value == "The lead engineer approves changes."


def test_exact_value_answer_needs_no_llm(engine, tmp_path):
    (tmp_path / "calc.txt").write_text(CALC_TEXT, encoding="utf-8")
    engine.ingest_file(tmp_path / "calc.txt")
    llm = FakeLLM()
    ans = engine.answer("What is the push-up distance?", llm_status(llm))
    assert ans.mode == "exact" and ans.text == "10.6 mm" and not llm.calls
    assert ans.fact.item == "9.(1)" and ans.passages[0].label == "calc.txt"
    by_symbol = engine.answer("What is σra?", llm_status(llm))
    assert by_symbol.mode == "exact" and by_symbol.text == "111 N/mm²"
    assert by_symbol.fact.formula == "σra = Pd / Pr x σr"


def test_unrelated_or_near_miss_questions_do_not_match_exactly(loaded_engine, tmp_path):
    (tmp_path / "calc.txt").write_text(CALC_TEXT, encoding="utf-8")
    loaded_engine.ingest_file(tmp_path / "calc.txt")
    for q in ["What is the push-up force?", "How many earned leave days do I get per year?",
              "What is the grip length?"]:
        assert loaded_engine.match_fact(q) is None, q


def test_word_guard_rejects_a_similar_name_with_a_different_meaning():
    [fact] = extract_facts([Page(1, "(1) Push-up distance, 10.6 mm")], "calc.pdf")
    assert question_covered("What is the push-up distance in mm?", fact)
    assert not question_covered("What is the push-up force?", fact)


def test_two_documents_with_different_values_are_not_answered_exactly(engine, tmp_path):
    (tmp_path / "a.txt").write_text("(1) Push-up distance, 10.6 mm", encoding="utf-8")
    (tmp_path / "b.txt").write_text("(1) Push-up distance, 12.0 mm", encoding="utf-8")
    engine.ingest_file(tmp_path / "a.txt")
    engine.ingest_file(tmp_path / "b.txt")
    assert engine.match_fact("What is the push-up distance?") is None


# ---- number check ------------------------------------------------------------ #
def test_numbers_missing_from_the_sources_are_flagged():
    passages = [Passage("Push-up distance, 10.6 mm. A= 1,118,355", "calc.pdf", 3, "", 0.9, "d")]
    answer = 'It is 10.6 mm ("Push-up distance, 10.6 mm") and the area is 1118355 [calc.pdf, p.3]; total 21.2 mm.'
    assert unverified_numbers(answer, passages) == ["21.2"]  # computed by the model, not in the document
    assert unverified_numbers("Item 9 says 10.6 mm", passages, question="What does item 9 say?") == []


def test_llm_answer_with_an_invented_number_gets_a_warning(engine, tmp_path):
    (tmp_path / "notes.txt").write_text("The push-up distance must be measured twice at 20 °C.", encoding="utf-8")
    engine.ingest_file(tmp_path / "notes.txt")
    llm = FakeLLM("Measure it twice at 25 °C [notes.txt]")
    ans = engine.answer("When must the push-up distance be measured?", llm_status(llm))
    assert ans.mode == "claude" and ans.unverified == ["25"]
